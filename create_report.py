"""Generate report.docx — v8 is the best model."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

ROOT = Path(__file__).resolve().parent
doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(1.5)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name  = "Times New Roman"
        run.font.size  = Pt(14 if level == 1 else 12)
        run.font.bold  = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_para(doc, text, size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.font.name   = "Times New Roman"
    run.font.size   = Pt(size)
    run.font.bold   = bold
    run.font.italic = italic
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


# ── Title ────────────────────────────────────────────────────────────────────
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("ПОЯСНИТЕЛЬНАЯ ЗАПИСКА")
r.font.name = "Times New Roman"; r.font.size = Pt(14); r.font.bold = True

s1 = doc.add_paragraph()
s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s1.add_run("Прогнозирование почасовой выработки ветроэлектростанции")
r.font.name = "Times New Roman"; r.font.size = Pt(13); r.font.bold = True

s2 = doc.add_paragraph()
s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s2.add_run("Хакатон «Hack the wind — claim the win», 2026 г.")
r.font.name = "Times New Roman"; r.font.size = Pt(12)

s3 = doc.add_paragraph()
s3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s3.add_run("Лучшая модель: LightGBM v8 — val MAE = 6,99 МВт (nMAE = 7,76 %)")
r.font.name = "Times New Roman"; r.font.size = Pt(12); r.font.italic = True

doc.add_paragraph()

# ── 1. Описание задачи ───────────────────────────────────────────────────────
add_heading(doc, "1. Описание задачи")
add_para(doc,
    "Задача хакатона — разработать математическую модель прогнозирования почасовой "
    "выработки ветроэлектростанции (ВЭС) установленной мощностью H_уст = 90,09 МВт "
    "(26 турбин Siemens Gamesa SG 3.4-132, координаты 46,83° с.ш. / 38,72° в.д., "
    "высота ротора 80 м). Входные данные — метеопрогноз: скорость и направление ветра "
    "на нескольких высотах, температура, давление, осадки."
)
add_para(doc, "Метрика качества — погрешность прогноза X, %:", bold=True)
add_para(doc, "    X = |S_факт − S_прогноз| / H_уст × 100 %")
add_para(doc,
    "Цель — минимизировать X. Победитель определяется как участник с наименьшей "
    "средней погрешностью на тестовой выборке."
)

# ── 2. Подход ────────────────────────────────────────────────────────────────
add_heading(doc, "2. Выбранный подход")
add_para(doc,
    "Решение построено на гомогенном LightGBM-ансамбле с гиперпараметрами, "
    "подобранными алгоритмом Optuna. Ключевые принципы:"
)
for item in [
    "Физически мотивированный инжиниринг признаков: эмпирическая кривая мощности, "
    "u/v-компоненты ветра, плотность воздуха, индекс обледенения и турбулентности.",
    "Отказ от лагов целевой переменной — устраняет накопление ошибок при авторегрессии "
    "и улучшает обобщение на новые данные.",
    "Оптимизация гиперпараметров через Optuna (200 проб, TPE). Результат сохранён "
    "в optuna_v8.json и используется повторно без перебора.",
    "Ансамбль из 15 членов LightGBM GBDT (5 конфигураций × 3 seed) — "
    "diversity по начальным условиям снижает дисперсию прогноза.",
    "Взвешенное объединение через scipy.optimize.minimize (SLSQP) — "
    "оптимальные веса на val-сплите вместо простого среднего.",
    "Финальное дообучение на полном датасете с best_iter × 1,15 итерациями.",
]:
    add_bullet(doc, item)

# ── 3. Архитектура ───────────────────────────────────────────────────────────
add_heading(doc, "3. Архитектура модели v8")

add_heading(doc, "3.1 Состав ансамбля", level=2)
add_para(doc, "Итоговая модель — ансамбль из 15 членов LightGBM GBDT:")

tbl = doc.add_table(rows=2, cols=3)
tbl.style = "Table Grid"
for cell, txt in zip(tbl.rows[0].cells, ["Тип", "Кол-во", "Конфигурации"]):
    cell.text = txt
    cell.paragraphs[0].runs[0].font.bold = True
    cell.paragraphs[0].runs[0].font.size = Pt(11)
    cell.paragraphs[0].runs[0].font.name = "Times New Roman"
for cell, val in zip(tbl.rows[1].cells,
                     ["LightGBM GBDT", "15", "5 лучших Optuna-конфигураций × 3 seed"]):
    cell.text = val
    cell.paragraphs[0].runs[0].font.size = Pt(11)
    cell.paragraphs[0].runs[0].font.name = "Times New Roman"
doc.add_paragraph()

add_heading(doc, "3.2 Инжиниринг признаков", level=2)
add_para(doc, "Из 19 исходных столбцов метеопрогноза формируется 251 признак:")
for item in [
    "Циклические календарные: sin/cos часа, месяца, дня года, дня недели.",
    "Ветровые компоненты: (u, v)-компоненты на 4 высотах; wind veer между уровнями.",
    "Физика ветра: v², v³ на каждой высоте; профиль сдвига скоростей; прокси hub-height.",
    "Термодинамика: плотность воздуха ρ = P/(R·T); ρ·v³; ρ·v³_hub.",
    "Эмпирическая кривая мощности: pc_pred (МВт) — восстановлена из обучающих данных "
    "методом квантильного сглаживания по скорости ветра.",
    "Лаги/лиды метео: сдвиги −1..−48 ч и +1..+6 ч по ключевым переменным.",
    "Rolling stats: mean/std за 3, 6, 12, 24 ч по скоростям ветра.",
    "Дополнительные: turbulence_index, icing_risk, pressure_change, precip_total.",
]:
    add_bullet(doc, item)

add_heading(doc, "3.3 Оптимизация гиперпараметров", level=2)
add_para(doc,
    "Гиперпараметры LightGBM подбирались алгоритмом Optuna (TPE, 200 проб) "
    "на хронологическом сплите train/val (последние 10 % обучающих данных). "
    "Оптимизируемые параметры: learning_rate, num_leaves, min_child_samples, "
    "reg_lambda, colsample_bytree, subsample. "
    "Лучшие 5 конфигураций сохраняются в optuna_v8.json и переиспользуются "
    "при каждом запуске train_v8.py."
)

add_heading(doc, "3.4 Взвешенный ансамбль", level=2)
add_para(doc,
    "Предсказания 15 членов объединяются взвешенным голосованием. Веса оптимизируются "
    "через scipy.optimize.minimize с ограничениями: все веса ≥ 0, сумма = 1, "
    "критерий — MAE на валидационной выборке. Это даёт снижение MAE по сравнению "
    "с равновзвешенным ансамблем: 7,16 → 6,99 МВт."
)

# ── 4. Результаты ─────────────────────────────────────────────────────────────
add_heading(doc, "4. Результаты на валидационной выборке")
add_para(doc,
    "Валидация = последние 10 % обучающих данных (~3 244 часа, хронологический сплит). "
    "Данные за 2026 год не использовались при обучении."
)

tbl2 = doc.add_table(rows=4, cols=3)
tbl2.style = "Table Grid"
for cell, txt in zip(tbl2.rows[0].cells, ["Конфигурация", "MAE, МВт", "nMAE, %"]):
    cell.text = txt
    cell.paragraphs[0].runs[0].font.bold = True
    cell.paragraphs[0].runs[0].font.size = Pt(11)
    cell.paragraphs[0].runs[0].font.name = "Times New Roman"
for i, (conf, mae, nmae) in enumerate([
    ("Базовая LightGBM, 7 членов (v3)",         "7,63", "8,47 %"),
    ("LightGBM Optuna v8, равные веса",          "7,16", "7,95 %"),
    ("LightGBM Optuna v8, оптим. веса (лучшая)", "6,99", "7,76 %"),
], 1):
    row = tbl2.rows[i].cells
    for cell, val in zip(row, [conf, mae, nmae]):
        cell.text = val
        cell.paragraphs[0].runs[0].font.size = Pt(11)
        cell.paragraphs[0].runs[0].font.name = "Times New Roman"
        if i == 3:
            cell.paragraphs[0].runs[0].font.bold = True
doc.add_paragraph()

add_para(doc,
    "Метрика nMAE % соответствует формуле погрешности из условия хакатона: "
    "X = |S_факт − S_прогноз| / H_уст × 100 %, H_уст = 90,09 МВт."
)

# ── 5. Воспроизводимость ─────────────────────────────────────────────────────
add_heading(doc, "5. Воспроизводимость")
add_para(doc, "Для воспроизведения результата необходимо:")
for i, step in enumerate([
    "Установить зависимости: pip install -r requirements.txt (Python 3.14).",
    "Подобрать гиперпараметры: python tune.py --trials 200  "
    "(или использовать готовый кэш optuna_v8.json — шаг можно пропустить).",
    "Обучить v8: python train_v8.py → models/model_v8.pkl.",
    "Сгенерировать предсказания: python predict.py --model models/model_v8.pkl "
    "--out predictions_v8.csv → файл с заголовком predict, 2126 строк.",
], 1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(f"{i}. {step}")
    r.font.name = "Times New Roman"; r.font.size = Pt(12)

add_para(doc,
    "Воспроизводимость обеспечена фиксацией SEED = 42 через set_global_seed "
    "(numpy, random, Python hash). При идентичных входных данных и версиях "
    "библиотек результат полностью воспроизводится."
)

out = ROOT / "report.docx"
doc.save(str(out))
print(f"Saved -> {out}")
